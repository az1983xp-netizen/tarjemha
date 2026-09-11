import os, sqlite3, uuid, shutil, time, threading, re
from pathlib import Path
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pypdf import PdfReader
from docx import Document
from docx.shared import Pt

BASE=Path(__file__).resolve().parent
UPLOADS=BASE/"uploads"; OUTPUTS=BASE/"outputs"
UPLOADS.mkdir(exist_ok=True); OUTPUTS.mkdir(exist_ok=True)
DB=BASE/"app.db"
OPENAI_API_KEY=os.getenv("OPENAI_API_KEY","")
OPENAI_MODEL=os.getenv("OPENAI_MODEL","gpt-5.6-luna")
app=FastAPI(title="Tarjemha",version="3.0.0")

def db():
    c=sqlite3.connect(DB); c.row_factory=sqlite3.Row; return c
with db() as c:
    c.execute("""CREATE TABLE IF NOT EXISTS orders(
      id TEXT PRIMARY KEY, filename TEXT, source TEXT, target TEXT, domain TEXT,
      pages INTEGER, price REAL, status TEXT, created_at REAL,
      original_path TEXT, text_path TEXT, output_path TEXT, error TEXT)""")

def estimate_price(pages, domain):
    m={"general":1,"academic":1.25,"legal":1.5,"medical":1.5,"business":1.15,"technical":1.3}
    return round(max(5,pages*2.5*m.get(domain,1)),2)

def extract(path):
    s=path.suffix.lower()
    if s==".txt": return path.read_text(encoding="utf-8",errors="ignore")
    if s==".docx": return "\n".join(p.text for p in Document(path).paragraphs)
    if s==".pdf":
        r=PdfReader(str(path)); return "\n".join((p.extract_text() or "") for p in r.pages)
    return ""

def page_count(path,text):
    if path.suffix.lower()==".pdf":
        try:return max(1,len(PdfReader(str(path)).pages))
        except:return 1
    return max(1, round(max(len(text),1)/2200))

def chunks(text, max_chars=12000):
    # Keep chunks on paragraph boundaries where possible.
    paras=text.split("\n")
    out=[]; cur=""
    for p in paras:
        if len(cur)+len(p)+1 > max_chars and cur:
            out.append(cur); cur=""
        cur += p+"\n"
    if cur.strip(): out.append(cur)
    return out

def translate_chunk(text, source, target, domain):
    if not OPENAI_API_KEY:
        raise RuntimeError("OPENAI_API_KEY is not configured.")
    from openai import OpenAI
    client=OpenAI(api_key=OPENAI_API_KEY)
    system=f"""You are a professional translator for {domain} documents.
Translate from {source} to {target}.
Preserve meaning, numbers, units, names, citations, equations, headings, and paragraph structure.
Do not summarize. Do not add commentary. Return only the translation.
If a technical term has an established target-language equivalent, use it consistently."""
    r=client.responses.create(
        model=OPENAI_MODEL,
        instructions=system,
        input=text,
        store=False
    )
    return r.output_text

def quality_check(original, translated, source, target, domain):
    if not OPENAI_API_KEY: return translated
    from openai import OpenAI
    client=OpenAI(api_key=OPENAI_API_KEY)
    prompt=f"""Review this {domain} translation from {source} to {target}.
Fix clear mistranslations, missing content, wrong numbers, names, units, or inconsistent terminology.
Do not summarize and do not explain changes. Return only the corrected translation.

ORIGINAL:
{original}

TRANSLATION:
{translated}"""
    r=client.responses.create(model=OPENAI_MODEL,instructions="You are a meticulous translation QA editor.",input=prompt,store=False)
    return r.output_text

def build_docx(source_path, translated, output_path):
    doc=Document()
    for block in translated.split("\n"):
        p=doc.add_paragraph(block)
        for run in p.runs: run.font.size=Pt(11)
    doc.save(output_path)

def process(order_id):
    try:
        with db() as c:
            row=c.execute("SELECT * FROM orders WHERE id=?",(order_id,)).fetchone()
            c.execute("UPDATE orders SET status=?,error=NULL WHERE id=?",("processing",order_id)); c.commit()
        text=Path(row["text_path"]).read_text(encoding="utf-8")
        pieces=chunks(text)
        translated_parts=[]
        for piece in pieces:
            translated_parts.append(translate_chunk(piece,row["source"],row["target"],row["domain"]))
        translated="\n".join(translated_parts)
        translated=quality_check(text,translated,row["source"],row["target"],row["domain"])
        outpath=OUTPUTS/f"{order_id}.docx"
        build_docx(Path(row["original_path"]),translated,outpath)
        with db() as c:
            c.execute("UPDATE orders SET status=?,output_path=? WHERE id=?",("completed",str(outpath),order_id)); c.commit()
    except Exception as e:
        with db() as c:
            c.execute("UPDATE orders SET status=?,error=? WHERE id=?",("failed",str(e)[:1000],order_id)); c.commit()

@app.get("/api/health")
def health():
    return {"status":"ok","version":"3.0.0","translation_configured":bool(OPENAI_API_KEY)}

@app.post("/api/orders")
async def create_order(file:UploadFile=File(...),source_language:str=Form(...),target_language:str=Form(...),domain:str=Form("general")):
    if source_language==target_language: raise HTTPException(400,"اختر لغتين مختلفتين.")
    suffix=Path(file.filename or "").suffix.lower()
    if suffix not in {".pdf",".docx",".txt"}: raise HTTPException(400,"الملفات المدعومة: PDF, DOCX, TXT.")
    oid=uuid.uuid4().hex[:10].upper()
    original=UPLOADS/f"{oid}{suffix}"
    with original.open("wb") as f: shutil.copyfileobj(file.file,f)
    text=extract(original)
    if not text.strip(): raise HTTPException(400,"لم أستطع استخراج النص. ملفات الصور الممسوحة تحتاج OCR.")
    tp=UPLOADS/f"{oid}.txt"; tp.write_text(text,encoding="utf-8")
    pg=page_count(original,text); price=estimate_price(pg,domain)
    with db() as c:
        c.execute("INSERT INTO orders VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                  (oid,file.filename,source_language,target_language,domain,pg,price,
                   "awaiting_payment",time.time(),str(original),str(tp),None,None)); c.commit()
    return {"id":oid,"filename":file.filename,"pages":pg,"price":price,"status":"awaiting_payment",
            "translation_configured":bool(OPENAI_API_KEY)}

@app.get("/api/orders/{oid}")
def order(oid:str):
    with db() as c: r=c.execute("SELECT id,filename,source,target,domain,pages,price,status,created_at,error FROM orders WHERE id=?",(oid,)).fetchone()
    if not r: raise HTTPException(404,"الطلب غير موجود.")
    return dict(r)

@app.post("/api/orders/{oid}/demo-pay")
def demo_pay(oid:str):
    with db() as c: r=c.execute("SELECT id FROM orders WHERE id=?",(oid,)).fetchone()
    if not r: raise HTTPException(404,"الطلب غير موجود.")
    if not OPENAI_API_KEY: raise HTTPException(503,"الترجمة غير مفعلة بعد. أضف مفتاح مزود الذكاء الاصطناعي.")
    with db() as c: c.execute("UPDATE orders SET status=? WHERE id=?",("paid",oid)); c.commit()
    threading.Thread(target=process,args=(oid,),daemon=True).start()
    return {"id":oid,"status":"processing"}

@app.get("/api/orders/{oid}/download")
def download(oid:str):
    with db() as c: r=c.execute("SELECT output_path,filename,status FROM orders WHERE id=?",(oid,)).fetchone()
    if not r or r["status"]!="completed": raise HTTPException(404,"الترجمة ليست جاهزة بعد.")
    return FileResponse(r["output_path"],filename=f"translated_{Path(r['filename']).stem}.docx",
                        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

@app.get("/")
def home(): return FileResponse(BASE.parent/"frontend"/"index.html")
app.mount("/assets",StaticFiles(directory=BASE.parent/"frontend"/"assets"),name="assets")
